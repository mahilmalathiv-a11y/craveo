import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def replace_text_frame_content(tf, lines, is_bullet=True):
    # Clear all but the first paragraph
    while len(tf.paragraphs) > 1:
        p = tf.paragraphs[-1]
        p_element = p._p
        p_element.getparent().remove(p_element)
        
    for idx, line in enumerate(lines):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        p.space_after = Pt(6)
        
        # Check if line contains a bullet point or prefix
        parts = line.split(":", 1)
        if len(parts) == 2 and not line.strip().startswith("•") and not line.strip().startswith("http"):
            bold_prefix = parts[0].strip()
            rest_text = parts[1]
            
            p.text = "" # Clear runs
            if is_bullet:
                run_bullet = p.add_run()
                run_bullet.text = "• "
                run_bullet.font.bold = True
                run_bullet.font.name = "Segoe UI"
                run_bullet.font.size = Pt(13)
                run_bullet.font.color.rgb = RGBColor(0, 69, 143)
                
            run1 = p.add_run()
            run1.text = bold_prefix + ": "
            run1.font.bold = True
            run1.font.name = "Segoe UI"
            run1.font.size = Pt(13)
            run1.font.color.rgb = RGBColor(0, 69, 143) # Deep Blue accent
            
            run2 = p.add_run()
            run2.text = rest_text
            run2.font.bold = False
            run2.font.name = "Segoe UI"
            run2.font.size = Pt(13)
            run2.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
        else:
            p.text = "" # Clear runs
            if is_bullet and not line.strip().startswith("•") and not line.endswith(":"):
                run_bullet = p.add_run()
                run_bullet.text = "• "
                run_bullet.font.bold = True
                run_bullet.font.name = "Segoe UI"
                run_bullet.font.size = Pt(13)
                run_bullet.font.color.rgb = RGBColor(0, 69, 143)
                
            run = p.add_run()
            run.text = line
            run.font.name = "Segoe UI"
            if line.endswith(":"):
                run.font.bold = True
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0, 69, 143)
            else:
                run.font.bold = False
                run.font.size = Pt(13)
                run.font.color.rgb = RGBColor(51, 65, 85)

def set_title_text(shape, new_title):
    tf = shape.text_frame
    tf.text = ""
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = new_title
    run.font.bold = True
    run.font.name = "Segoe UI"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 69, 143)

def main():
    path = "D:\\INTERN PROJECT\\Internship Project Template.pptx"
    out_path = "D:\\INTERN PROJECT\\Craveo_Internship_Presentation.pptx"
    
    prs = Presentation(path)
    
    # ------------------------------------------
    # Slide 1: Cover
    # ------------------------------------------
    # We do not modify the cover image slide.
    
    # ------------------------------------------
    # Slide 2: Title Slide
    # ------------------------------------------
    s2 = prs.slides[1]
    title_box = s2.shapes[1]
    tf2 = title_box.text_frame
    tf2.text = ""
    p2 = tf2.paragraphs[0]
    run2_title = p2.add_run()
    run2_title.text = "CRAVEO\n"
    run2_title.font.bold = True
    run2_title.font.name = "Segoe UI"
    run2_title.font.size = Pt(36)
    run2_title.font.color.rgb = RGBColor(0, 69, 143)
    
    p2_sub = tf2.add_paragraph()
    run2_sub = p2_sub.add_run()
    run2_sub.text = "AI-Powered Food Ordering & Delivery ETA Engine\n\nFull-Stack Logistics System with XGBoost Regressor\n"
    run2_sub.font.bold = False
    run2_sub.font.name = "Segoe UI"
    run2_sub.font.size = Pt(18)
    run2_sub.font.color.rgb = RGBColor(51, 65, 85)
    
    p2_pres = tf2.add_paragraph()
    run2_pres = p2_pres.add_run()
    run2_pres.text = "Presented by: Developer Team"
    run2_pres.font.bold = True
    run2_pres.font.name = "Segoe UI"
    run2_pres.font.size = Pt(14)
    run2_pres.font.color.rgb = RGBColor(0, 69, 143)

    # ------------------------------------------
    # Slide 3: Slide 3 Overview
    # ------------------------------------------
    s3 = prs.slides[2]
    set_title_text(s3.shapes[2], "1. Project Overview & Objectives")
    s3_bullets = [
        "The Business Problem: Uncertainty in delivery times leads to customer frustration, kitchen queue congestion, and driver dispatch inefficiencies.",
        "Target Objective: Build an AI-powered Django & MySQL system that predicts food delivery durations in real-time at checkout.",
        "Core Solution: A trained XGBoost Regressor combined with responsive multi-role dashboards for customers, owners, and drivers."
    ]
    replace_text_frame_content(s3.shapes[0].text_frame, s3_bullets)

    # ------------------------------------------
    # Slide 4: Slide 4 Architecture
    # ------------------------------------------
    s4 = prs.slides[3]
    set_title_text(s4.shapes[2], "2. System Architecture & Relational Design")
    s4_bullets = [
        "Responsive Client Layer: Role-aware dashboards built with Bootstrap 5 and JavaScript for multi-role workflows.",
        "Django Application Core: Handles MVC routing, session cart checkout, and secure MySQL relational transactions.",
        "AI Inference Service: Standard scaler and serialized XGBoost model that executes estimations in under 10ms.",
        "Database Architecture: Relational MySQL schema with 8 custom tables and built-in SQL error analytics."
    ]
    replace_text_frame_content(s4.shapes[0].text_frame, s4_bullets)

    # ------------------------------------------
    # Slide 5: Preprocessing & Feature Engineering
    # ------------------------------------------
    s5 = prs.slides[4]
    set_title_text(s5.shapes[2], "3. Data Preprocessing & IQR Capping")
    s5_bullets_3 = [
        "Deduplication: Removes redundant order logs to prevent model overfitting.",
        "Median Imputation: Fills missing driver ratings with median scores.",
        "IQR Capping: Caps extreme times and distances using Q3 + 1.5 * IQR."
    ]
    replace_text_frame_content(s5.shapes[0].text_frame, s5_bullets_3)
    
    set_title_text(s5.shapes[3], "4. Feature Engineering & Selection")
    s5_bullets_4 = [
        "Haversine Distance: Spherical distance computed locally between GPS points.",
        "Spatio-Temporal Variables: Hour, day of week, and rush hour classifications.",
        "Operational Context: Kitchen active queue load and checkout order volume."
    ]
    replace_text_frame_content(s5.shapes[4].text_frame, s5_bullets_4)

    # ------------------------------------------
    # Slide 6: Model Selection & XGBoost
    # ------------------------------------------
    s6 = prs.slides[5]
    set_title_text(s6.shapes[2], "5. Model Selection & Rationale")
    
    s6_benchmark = [
        "Comparative Model Benchmarking:",
        "Linear Regression: R² = 91.41% | MAE = 4.47 mins",
        "Random Forest Regressor: R² = 95.71% | MAE = 3.01 mins",
        "XGBoost Regressor (Selected): R² = 95.93% | MAE = 2.79 mins"
    ]
    replace_text_frame_content(s6.shapes[0].text_frame, s6_benchmark, is_bullet=False)
    
    s6_xgboost = [
        "Why XGBoost is the Optimal Choice:",
        "Sequentially minimizes mean squared error loss by learning from previous tree residuals.",
        "Regularization controls model complexity to prevent overfitting on noisy logistics log data.",
        "Evaluates complex multi-modal variables in under 10 milliseconds."
    ]
    replace_text_frame_content(s6.shapes[3].text_frame, s6_xgboost, is_bullet=False)

    # ------------------------------------------
    # Slide 7: Validation & Screenshots
    # ------------------------------------------
    s7 = prs.slides[6]
    set_title_text(s7.shapes[2], "6. Project Validation & Dashboards")
    s7_bullets = [
        "System Validation & Visual Interfaces:",
        "Prediction Accuracy: 86.00% of estimations fall within ±5 minutes of actual times.",
        "Operational Stepper: Progress tracker (Pending -> Preparing -> Dispatch -> Delivered).",
        "Admin Performance Console: Direct Chart.js dashboards monitoring active queues.",
        "Power BI Integration: Relational database CSV exporter streaming live metadata."
    ]
    replace_text_frame_content(s7.shapes[0].text_frame, s7_bullets, is_bullet=False)

    # ------------------------------------------
    # Slide 8: Thank You
    # ------------------------------------------
    s8 = prs.slides[7]
    tf8 = s8.shapes[0].text_frame
    tf8.text = ""
    p8 = tf8.paragraphs[0]
    p8.alignment = PP_ALIGN.CENTER
    run8 = p8.add_run()
    run8.text = "Thank You\n\nQuestions & Answers"
    run8.font.bold = True
    run8.font.name = "Segoe UI"
    run8.font.size = Pt(28)
    run8.font.color.rgb = RGBColor(0, 69, 143)

    prs.save(out_path)
    print(f"Template populated and saved successfully to: {out_path}")

if __name__ == "__main__":
    main()
