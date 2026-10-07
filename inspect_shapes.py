from pptx import Presentation

def main():
    path = "D:\\INTERN PROJECT\\Internship Project Template.pptx"
    prs = Presentation(path)
    for idx, slide in enumerate(prs.slides):
        print(f"\n================= SLIDE {idx + 1} =================")
        for s_idx, shape in enumerate(slide.shapes):
            print(f"Shape {s_idx}: Name='{shape.name}', ID={shape.shape_id}, Type={shape.shape_type}")
            if shape.has_text_frame:
                print(f"  Text: {repr(shape.text)}")
                for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                    runs = [run.text for run in paragraph.runs]
                    print(f"    Paragraph {p_idx}: runs={runs}")

if __name__ == "__main__":
    main()
