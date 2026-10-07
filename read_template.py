from pptx import Presentation

def main():
    path = "D:\\INTERN PROJECT\\Internship Project Template.pptx"
    try:
        prs = Presentation(path)
        print(f"Successfully loaded presentation. Number of slides: {len(prs.slides)}")
        for idx, slide in enumerate(prs.slides):
            print(f"\n--- SLIDE {idx + 1} ---")
            
            # Print slide title if it exists
            if slide.shapes.title:
                print(f"Title: {slide.shapes.title.text}")
            else:
                print("Title: [No Title Shape]")
                
            # Iterate through shapes and look for text
            for shape in slide.shapes:
                if shape.has_text_frame:
                    # Skip the title shape since we already printed it
                    if shape == slide.shapes.title:
                        continue
                    text = shape.text.strip()
                    if text:
                        print(f"  Shape [{shape.name}]:")
                        print(f"    {text}")
    except Exception as e:
        print(f"Error loading template: {e}")

if __name__ == "__main__":
    main()
